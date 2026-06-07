from __future__ import annotations

import io
from pathlib import Path

from schema.models import PreviewView


def export_markdown(preview: PreviewView) -> str:
    lines = [f"# {preview.paper_title or '论文'}", ""]
    current_section = None
    for span in preview.spans:
        if span.section_id != current_section:
            current_section = span.section_id
            lines.append("")
        marker = ""
        if span.highlight == "accepted":
            marker = " ✓"
        elif span.highlight == "custom":
            marker = " ✎"
        lines.append(f"{span.text}{marker}")
    return "\n".join(lines).strip() + "\n"


def export_docx_bytes(preview: PreviewView) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for docx export") from exc

    doc = Document()
    doc.add_heading(preview.paper_title or "论文", level=0)
    for span in preview.spans:
        doc.add_paragraph(span.text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_pdf_bytes(preview: PreviewView) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError("reportlab is required for pdf export") from exc

    font_name = "STSong-Light"
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    except Exception:
        font_name = "Helvetica"

    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 50
    line_height = 16

    page.setFont(font_name, 14)
    page.drawString(50, y, preview.paper_title or "论文")
    y -= line_height * 2
    page.setFont(font_name, 11)

    for span in preview.spans:
        text = span.text
        if not text:
            continue
        for line in _wrap_pdf_text(text, font_name, 11, width - 100):
            if y < 60:
                page.showPage()
                page.setFont(font_name, 11)
                y = height - 50
            page.drawString(50, y, line)
            y -= line_height

    page.save()
    return buffer.getvalue()


def _wrap_pdf_text(text: str, font_name: str, font_size: int, max_width: float) -> list[str]:
    try:
        from reportlab.pdfbase import pdfmetrics
    except ImportError as exc:
        raise RuntimeError("reportlab is required for pdf export") from exc

    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if current and pdfmetrics.stringWidth(candidate, font_name, font_size) > max_width:
            lines.append(current)
            current = char
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines or [""]


def write_export_file(preview: PreviewView, fmt: str, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        dest.write_bytes(export_docx_bytes(preview))
    elif fmt == "md":
        dest.write_text(export_markdown(preview), encoding="utf-8")
    elif fmt == "pdf":
        dest.write_bytes(export_pdf_bytes(preview))
    else:
        raise ValueError(f"unsupported export format: {fmt}")
    return dest
